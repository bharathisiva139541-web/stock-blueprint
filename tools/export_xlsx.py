"""Export a downloaded copy of the Stock Blueprint sheet (.xlsx) into the
website's data file  web/data/dashboard.js  (sets window.SB_DATA).

    python tools/export_xlsx.py "Copy of Stock Blueprint.xlsx"

Only market / signal tabs are exported. Personal data is never exported,
because the website is public: Portfolio, Risk Manager, Control Panel,
Trade Log, Trade Dashboard, and 'Entry' alerts (those carry the user's own
entry prices / open positions).
"""
import datetime as dt
import json
import math
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'web', 'data', 'dashboard.js')

ERRORS = ('#N/A', '#DIV/0!', '#REF!', '#VALUE!', '#ERROR!', '#NAME?', '#NUM!', '#NULL!', 'Loading...')

# Main tab: (output key, header name in row 2, kind)
#   kind: str | num | pct (already a percent number) | frac (fraction -> x100) | date | bool
MAIN_COLS = [
    ('sym', 'SYMBOL', 'str'),
    ('name', 'StockName', 'str'),
    ('sector', 'Sector', 'str'),
    ('industry', 'Industry', 'str'),
    ('ind', 'Basic Industry', 'str'),
    ('idx', 'INDICES', 'str'),
    ('rdate', 'Result Date', 'date'),
    ('done', 'Checkbox', 'bool'),
    ('mcap', 'MCAP(CR)', 'num'),
    ('price', 'MKT Price', 'num'),
    ('chg', '% Change', 'pct'),
    ('d10', '10 day%', 'frac'),
    ('d30', '30Day%', 'frac'),
    ('hi52', '52W High', 'num'),
    ('down52', '% down', 'frac'),
    ('pe', 'PE Ratio', 'num'),
    ('ipe', 'Industry PE', 'num'),
    ('pevs', 'PE vs Industry PE', 'num'),
    ('ttmEps', 'TTM EPS', 'num'),
    ('lfyEps', 'Last FY EPS', 'num'),
    ('epsg', 'TTM EPS Growth', 'num'),
    ('peg', 'TTM PEG Ratio', 'num'),
    ('tgt', 'TargetPrice', 'num'),
    ('sig', 'TTM Signal', 'str'),
    ('up', 'Upside Potential', 'frac'),
    ('sig2', 'Signal', 'str'),
    ('qrev', 'QoQ[Revenue Growth%}', 'pct'),
    ('qnp', 'QoQ[NP Growth %]', 'pct'),
    ('yrev', 'YoY[Revenue Growth %]', 'pct'),
    ('ynp', 'YoY[NP Growth %]', 'pct'),
    ('promo', 'Promotor Holding', 'str'),
    ('fii', 'FII Activity', 'str'),
    ('opmy', 'OPM YoY%', 'str'),
    ('opmq', 'OPM QoQ%', 'str'),
    ('pat', 'Pattern', 'str'),
    ('fvg', 'FVG', 'str'),
    ('e10', '10 EMA', 'str'),
    ('e21', '21 EMA', 'str'),
    ('zone', 'Zone', 'str'),
    ('volchg', '%vol Change', 'frac'),
    ('ipo', 'IPO Level', 'str'),
]


# Alert Log columns that may be published (no personal position data)
ALERT_COLS = ('Date', 'Symbol', 'Company', 'Alert Type', 'Alert Level', 'Trigger Price',
              'Pattern @ Alert', 'FVG Tags @ Alert', 'Current Price', '% Since Alert',
              'Days Since', 'Outcome')
PUBLIC_ALERT_TYPES = ('Breakout', 'Price Fall')


def clean(v, kind='auto'):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == '' or s in ERRORS or s.startswith('#'):
            return None
        if kind in ('num', 'pct', 'frac'):
            try:
                v = float(s.replace(',', ''))
            except ValueError:
                return None
        else:
            return text_date(s) or s
    if isinstance(v, bool):
        return v if kind in ('bool', 'auto') else None
    if isinstance(v, (dt.datetime, dt.date)):
        if kind in ('num', 'pct', 'frac'):
            return None
        if isinstance(v, dt.datetime) and (v.hour or v.minute):
            return v.strftime('%Y-%m-%d %H:%M')
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)):
        if kind == 'frac':
            v = v * 100
        if kind in ('str', 'date'):
            return str(v)
        return round(float(v), 4)
    return str(v)


TEXT_DATES = ('%d-%b-%Y %H:%M', '%d-%b-%Y', '%d-%b-%y', '%d/%m/%Y %H:%M', '%d/%m/%Y')


def text_date(s):
    """'21-Aug-2026', '01/10/2026 20:23' ... -> ISO text (sorts correctly)."""
    if not (8 <= len(s) <= 17 and s[:2].isdigit()):
        return None
    for f in TEXT_DATES:
        try:
            d = dt.datetime.strptime(s, f)
        except ValueError:
            continue
        return d.strftime('%Y-%m-%d %H:%M' if '%H' in f else '%Y-%m-%d')
    return None


def norm(h):
    return ' '.join(str(h or '').split()).lower()


class Rows(list):
    """Cell values row by row, plus .fmt: the matching number formats."""
    fmt = ()


def rows_of(ws):
    # read-only mode returns short (or empty) tuples for sparse rows: pad to 1
    out, fmt = Rows(), []
    for r in ws.iter_rows():
        vals = [c.value for c in r] or [None]
        out.append(vals)
        fmt.append([getattr(c, 'number_format', 'General') for c in r] or ['General'])
    out.fmt = fmt
    return out


def is_frac_format(f):
    """True for formats such as 0.00% that show a fraction as a percent
    (a quoted "%" is a literal sign: the value is already a percent)."""
    f = re.sub(r'"[^"]*"', '', str(f or ''))
    f = re.sub(r'\\.', '', f)
    return '%' in f


def col_kind(rows, idxs, j):
    """Infer frac/auto for column j from its first numeric data cell
    (idxs = the table's own data row numbers, so nothing below it counts)."""
    for i in idxs[:200]:
        r = rows[i]
        if j < len(r) and isinstance(r[j], (int, float)) and not isinstance(r[j], bool):
            fr = rows.fmt[i] if getattr(rows, 'fmt', None) and i < len(rows.fmt) else []
            return 'frac' if j < len(fr) and is_frac_format(fr[j]) else 'auto'
    return 'auto'


def table(rows, first_header, start=0, stop_blank=True, kinds=None, max_rows=None):
    """Find the row whose first cell is `first_header` (at or after `start`),
    read headers from it and data rows below until the first blank first cell."""
    for i in range(start, len(rows)):
        if rows[i] and norm(rows[i][0]) == norm(first_header):
            hdr = [str(h).strip() if h is not None else '' for h in rows[i]]
            while hdr and not hdr[-1]:
                hdr.pop()
            idxs = []
            for k in range(i + 1, len(rows)):
                r = rows[k]
                if not r or r[0] in (None, ''):
                    if stop_blank:
                        break
                    continue
                idxs.append(k)
                if max_rows and len(idxs) >= max_rows:
                    break
            kinds = kinds or {}
            ks = [kinds.get(h) or col_kind(rows, idxs, j) for j, h in enumerate(hdr)]
            out = [[clean(rows[k][j] if j < len(rows[k]) else None, ks[j]) for j in range(len(hdr))] for k in idxs]
            return {'cols': hdr, 'rows': out}
    return {'cols': [], 'rows': []}


def note_table(rows, first_header, **kw):
    """table() plus the tab's first cell as its note. Code.gs ab_site_noteTable_ does the same."""
    t = table(rows, first_header, **kw)
    first = rows[0][0] if rows and rows[0] and rows[0][0] is not None else ''
    return {'note': str(first).strip() if first != '' else '', 'cols': t['cols'], 'rows': t['rows']}


IPO_SECTIONS = {'open now': 'open', 'closed - awaiting listing': 'closed', 'upcoming': 'upcoming'}


def ipo_tracker(rows):
    """The IPO Tracker tab (filled from NSE by the PC bridge) -> {note, warning, open, closed, upcoming}
    (each {cols, rows}; the "None" row dropped). Code.gs ab_site_ipoTracker_ does the same."""
    out, found = {'note': '', 'warning': ''}, {}
    for i, r in enumerate(rows):
        t = str(r[0]).strip() if r and r[0] is not None else ''
        if i == 1 and t.startswith('Updated '):
            out['note'] = t
        if i < 4 and t.startswith('Could not read'):
            out['warning'] = t
        m = re.match(r'^(.+) \((\d+)\)$', t)
        k = IPO_SECTIONS.get(m.group(1).lower()) if m else None
        if k and k not in found:
            tb = table(rows, 'Company', start=i + 1)
            tb['rows'] = [x for x in tb['rows'] if not (x[0] == 'None' and all(v is None for v in x[1:]))]
            found[k] = tb
    for k in ('open', 'closed', 'upcoming'):
        out[k] = found.get(k, {'cols': [], 'rows': []})
    return out


def export(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = {'generated': dt.datetime.now().strftime('%Y-%m-%d %H:%M'), 'source': os.path.basename(path)}

    # Main
    m = rows_of(wb['Main'])
    hdr = [norm(h) for h in m[1]]
    idx = {}
    for key, name, kind in MAIN_COLS:
        try:
            idx[key] = (hdr.index(norm(name)), kind)
        except ValueError:
            print('  ! Main column not found:', name)
    sym_i = idx['sym'][0]
    stocks = []
    for r in m[2:]:
        if sym_i >= len(r) or r[sym_i] in (None, ''):
            continue
        stocks.append([clean(r[i] if i < len(r) else None, k) for (i, k) in idx.values()])
    data['main'] = {'cols': list(idx.keys()), 'rows': stocks}

    # Symbols that are (or were) the user's own trades: any Main row with an
    # Entry Price or Stoploss typed in, plus every Trade Log symbol. Their alert
    # rows are never published, because the alert levels and outcomes of a
    # trade reveal the position.
    private_syms = set()
    for name in ('entry price', 'stoploss'):
        if name in hdr:
            ci = hdr.index(name)
            for r in m[2:]:
                v = r[ci] if ci < len(r) else None
                if sym_i < len(r) and r[sym_i] and v not in (None, '') and clean(v, 'num'):
                    private_syms.add(str(r[sym_i]).strip().upper())
    if 'Trade Log' in wb.sheetnames:
        for r in rows_of(wb['Trade Log'])[1:]:
            if len(r) > 1 and r[1]:
                private_syms.add(str(r[1]).strip().upper())
    # holdings (Portfolio 'Symbol' column) and the Risk Manager's open trades
    for tab in ('Portfolio', 'Risk Manager'):
        if tab not in wb.sheetnames:
            continue
        tr = rows_of(wb[tab])
        for i, r in enumerate(tr):
            if r and norm(r[0]) == 'symbol':
                for r2 in tr[i + 1:]:
                    v = r2[0] if r2 else None
                    if isinstance(v, str) and re.fullmatch(r'[A-Z0-9&_-]{1,20}', v.strip().upper()):
                        private_syms.add(v.strip().upper())

    # Results season
    data['resultsTracker'] = table(rows_of(wb['Results Tracker']), 'Symbol', stop_blank=False)
    rp = rows_of(wb['Results Performance'])
    perf = {'note': clean(rp[1][0]) if len(rp) > 1 else None, 'seasons': [], 'verdicts': [], 'lines': []}
    for i in range(len(rp)):
        if rp[i] and norm(rp[i][0]) == 'season':
            j = i + 1
            while j < len(rp) and norm(rp[j][0]) != 'symbol':
                r = rp[j]
                if r and r[0] not in (None, '') and len(r) > 5 and r[1] is not None and not str(r[0]).startswith(('Best', 'Worst')):
                    perf['seasons'].append({'season': clean(r[0]), 'tracked': clean(r[1], 'num'),
                                            'avg': clean(r[2], 'frac'), 'median': clean(r[3], 'frac'),
                                            'above': clean(r[4]), 'soldOff': clean(r[5], 'num')})
                elif r and r[0]:
                    perf['lines'].append(clean(r[0]))
                if r and len(r) > 11 and r[7] not in (None, ''):
                    perf['verdicts'].append({'verdict': clean(r[7]), 'stocks': clean(r[8], 'num'),
                                             'avg': clean(r[9], 'frac'), 'positive': clean(r[10], 'frac'),
                                             'day': clean(r[11], 'frac')})
                j += 1
            break
    perf['table'] = table(rp, 'Symbol')
    data['resultsPerf'] = perf
    # "as of" = when the sheet last refreshed (from the Results Performance note)
    m_ = re.search(r'last refreshed (\d{2}/\d{2}/\d{4} \d{2}:\d{2})', perf['note'] or '')
    if m_:
        data['asOf'] = text_date(m_.group(1))

    # Signals
    pd_ = rows_of(wb['Pattern Dashboard'])
    kp = {}
    m_ = re.search(r'refreshed (\d{2}/\d{2}/\d{4} \d{2}:\d{2})', ' '.join(str(c) for r in pd_[:3] for c in r if c))
    for i in range(len(pd_)):
        if pd_[i] and norm(pd_[i][0]) == 'stocks scanned':
            kp = {str(h): clean(v) for h, v in zip(pd_[i], pd_[i + 1]) if h}
            break
    sigk = {k: 'frac' for k in ('Avg % Change', '% Advancing', 'Avg 10-day %', 'Avg 30-day %')}
    clk = {'Avg % Change': 'frac'}
    stk = {k: 'frac' for k in ('% Change', '10-day %', '30-day %')}
    data['patterns'] = {
        'asOf': text_date(m_.group(1)) if m_ else None,
        'kpis': kp,
        'signals': table(pd_, 'Type', kinds=sigk),
        'clusters': table(pd_, 'Basic Industry', kinds=clk),
        'stocks': table(pd_, 'Symbol', kinds=stk),
    }
    data['fvg'] = table(rows_of(wb['UnmitigatedFVG']), 'Symbol', stop_blank=False)
    data['gapup'] = table(rows_of(wb['GapUpFVG']), 'Symbol', stop_blank=False)
    data['watchLog'] = table(rows_of(wb['Signal Watchlist Log']), 'Date Flagged', stop_blank=False)
    if 'FVG Backtest' in wb.sheetnames:
        data['fvgBacktest'] = table(rows_of(wb['FVG Backtest']), 'Flagged On', stop_blank=False)
    if 'Mitigated FVG Log' in wb.sheetnames:
        data['mitigatedLog'] = table(rows_of(wb['Mitigated FVG Log']), 'Date', stop_blank=False)

    # Alerts: watch-level alerts only. 'Entry' alerts are the user's real
    # positions (alert level = their entry price), and the Trade Log / Trade
    # Dashboard are their personal trading results - never exported.
    al = table(rows_of(wb['Alert Log']), 'Date', stop_blank=False)
    keep = [c for c in al['cols'] if c in ALERT_COLS]
    ti = al['cols'].index('Alert Type') if 'Alert Type' in al['cols'] else None
    si = al['cols'].index('Symbol') if 'Symbol' in al['cols'] else None
    rows = [r for r in al['rows'] if ti is not None and r[ti] in PUBLIC_ALERT_TYPES
            and si is not None and str(r[si] or '').strip().upper() not in private_syms]
    data['alerts'] = {'cols': keep, 'rows': [[r[al['cols'].index(c)] for c in keep] for r in rows],
                      'note': 'watch alerts only; symbols with your own entry/stop-loss or a Trade Log row are left out'}
    print('  alerts: kept', len(rows), 'of', len(al['rows']), '- left out', len(private_syms), 'trade symbols')

    # Markets
    ix = table(rows_of(wb['Indextracker']), 'NSE Symbol', stop_blank=False)
    seen, rows = set(), []
    for r in ix['rows']:            # drop label rows and the repeated block
        if r[0] in seen or not any(isinstance(v, (int, float)) for v in r[1:]):
            continue
        seen.add(r[0])
        rows.append(r)
    data['indices'] = {'cols': ix['cols'], 'rows': rows}
    data['ipos'] = table(rows_of(wb['IPOs']), 'NSE SYMBOL', stop_blank=False)
    data['newIpos'] = table(rows_of(wb['New IPOs']), 'COMPANY NAME', stop_blank=False, max_rows=300)
    data['priceBands'] = table(rows_of(wb['Price Band Changes']), 'Report Date', stop_blank=False)
    if 'IPO Tracker' in wb.sheetnames:
        data['ipoTracker'] = ipo_tracker(rows_of(wb['IPO Tracker']))
    # NSE end-of-day (PC bridge): market statistics, 52-week highs, Delivery %
    if 'Market Statistics' in wb.sheetnames:
        data['marketStats'] = note_table(rows_of(wb['Market Statistics']), 'Date', stop_blank=False)
    if 'Market Day' in wb.sheetnames:
        data['marketDay'] = note_table(rows_of(wb['Market Day']), 'Symbol', stop_blank=False)
    if '52 Week High' in wb.sheetnames:
        data['high52'] = note_table(rows_of(wb['52 Week High']), 'Symbol', stop_blank=False)
    if 'Delivery %' in wb.sheetnames:
        data['delivery'] = table(rows_of(wb['Delivery %']), 'Symbol', stop_blank=False)

    export_history(wb, data.get('asOf'))
    export_history60(wb)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    js = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('// Generated by tools/export_xlsx.py - do not edit by hand\nwindow.SB_DATA = ' + js + ';\n')
    print('wrote', os.path.normpath(OUT), f'{len(js) / 1024:.0f} KB,', len(stocks), 'stocks')


def export_history(wb, as_of):
    """History30 daily candles (public market data) + the sheet's own pattern /
    FVG tags (PatternData W / X) so the website can compare its JS port with
    the sheet. Written to web/data/history.js as window.SB_HISTORY."""
    if 'History30' not in wb.sheetnames:
        print('  ! no History30 tab - history.js not written')
        return
    rows = wb['History30'].iter_rows(values_only=True)
    hdr = [norm(h) for h in next(rows)]
    ci = {k: hdr.index(k) for k in ('symbol', 'date', 'open', 'high', 'low', 'close', 'volume')}
    dates, didx, bars = [], {}, {}
    for r in rows:
        sym = r[ci['symbol']] if ci['symbol'] < len(r) else None
        if not sym:
            continue
        d = clean(r[ci['date']], 'date')
        d = text_date(d) if d and '/' in d else d
        if not d:
            continue
        d = d[:10]
        if d not in didx:
            didx[d] = len(dates)
            dates.append(d)
        vals = [clean(r[ci[k]], 'num') for k in ('open', 'high', 'low', 'close', 'volume')]
        bars.setdefault(str(sym).strip().upper(), []).append([didx[d]] + vals)
    # re-index dates in calendar order
    order = sorted(range(len(dates)), key=lambda i: dates[i])
    remap = {old: new for new, old in enumerate(order)}
    dates = [dates[i] for i in order]
    for sym in bars:
        bars[sym] = sorted(([remap[b[0]]] + b[1:] for b in bars[sym]), key=lambda b: b[0])
    tags = {}
    if 'PatternData' in wb.sheetnames:
        pr = wb['PatternData'].iter_rows(values_only=True)
        ph = [norm(h) for h in next(pr)]
        def col(name, default):
            return ph.index(name) if name in ph else default
        idx = [col('pattern', 22), col('long fvg touch', 23), col('10 ema support', 24),
               col('atr14 (ema)', 25), col('21 ema support', 26)]
        for r in pr:
            sym = r[3] if len(r) > 3 else None
            if not sym:
                continue
            vals = [r[i] if i < len(r) else None for i in idx]
            # [pattern W, fvg X, 10 EMA Y, ATR14 Z, 21 EMA AA]
            tags[str(sym).strip().upper()] = [v if isinstance(v, (int, float)) and not isinstance(v, bool) else (v or '') for v in vals]
    out = {'asOf': as_of, 'source': 'History30 tab (GOOGLEFINANCE daily candles)', 'timeframe': '1d',
           'dates': dates, 'cols': ['d', 'o', 'h', 'l', 'c', 'v'], 'bars': bars, 'sheetTags': tags}
    path = os.path.join(os.path.dirname(OUT), 'history.js')
    js = json.dumps(out, separators=(',', ':'))
    with open(path, 'w', encoding='utf-8') as f:
        f.write('// Generated by tools/export_xlsx.py - do not edit by hand\nwindow.SB_HISTORY = ' + js + ';\n')
    print('wrote', os.path.normpath(path), f'{len(js) / 1024:.0f} KB,', len(bars), 'symbols,', len(dates), 'dates')


# ---------------------------------------------------------------- History60 (NSE candles via the PC bridge)
# The same reading as Code.gs ab_h60Grid_ / ab_h60Adjust_ / ab_h60Older_ / ab_site_history60_ - keep both in step.
H60_LINK = '(linked)'
_MONTHS = ('jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec')


def h60_iso(v):
    """A day cell -> 'yyyy-mm-dd' or '' (Code.gs ab_h60Iso_)."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    s = '' if v is None else str(v).strip()
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return '%s-%s-%s' % m.groups()
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        return '%s-%02d-%02d' % (m.group(3), int(m.group(2)), int(m.group(1)))
    m = re.match(r'^(\d{1,2})[-/ ]([A-Za-z]{3})[a-z]*[-/ ](\d{4})$', s)
    if m and m.group(2).lower() in _MONTHS:
        return '%s-%02d-%02d' % (m.group(3), _MONTHS.index(m.group(2).lower()) + 1, int(m.group(1)))
    return ''


def _h60_num(x):
    try:
        return float(x)
    except ValueError:
        return float('nan')


def h60_cell(v):
    """'o h l c v p' -> [o, h, l, c, v, p] (v / p nan when missing) or None."""
    if not isinstance(v, str) or not v.strip():
        return None
    a = [_h60_num(x) for x in v.split()]
    if len(a) < 4 or not (a[0] > 0 and a[1] > 0 and a[2] > 0 and a[3] > 0):
        return None
    nan = float('nan')
    return [a[0], a[1], a[2], a[3], a[4] if len(a) > 4 and a[4] >= 0 else nan, a[5] if len(a) > 5 and a[5] > 0 else nan]


def h60_grid(rows):
    out = {'days': [], 'linked': [], 'syms': [], 'cells': {}}
    h = next((i for i, r in enumerate(rows[:5]) if r and str(r[0]).strip() == 'Symbol'), -1)
    if h == -1:
        return out
    cols, seen = [], set()
    for j, v in enumerate(rows[h]):
        d = h60_iso(v) if j else ''
        if d and d not in seen:
            seen.add(d)
            cols.append((j, d))
    cols.sort(key=lambda c: c[1])
    nxt = rows[h + 1] if h + 1 < len(rows) else None
    link = nxt if nxt and str(nxt[0]).strip() == H60_LINK else None
    out['days'] = [c[1] for c in cols]
    out['linked'] = [bool(link) and j < len(link) and str(link[j]).strip() == 'Y' for j, _ in cols]
    for r in rows[h + (2 if link else 1):]:
        sym = ('' if not r or r[0] is None else str(r[0])).strip().upper()
        if not sym or sym in out['cells']:
            continue
        out['syms'].append(sym)
        out['cells'][sym] = [h60_cell(r[j] if j < len(r) else None) for j, _ in cols]
    return out


def _r2(x):
    return math.floor(x * 100 + 0.5) / 100


def _vol(v, f):
    return v if math.isnan(v) else math.floor(v / f + 0.5)


def h60_adjust(g, sym):
    cells = g['cells'].get(sym) or []
    out, f, newer = [], 1, -1
    for i in range(len(cells) - 1, -1, -1):
        c = cells[i]
        if not c:
            continue
        if newer != -1:
            gap = any(not g['linked'][k] for k in range(i + 1, newer + 1))
            p = cells[newer][5]
            if not gap and p > 0 and abs(p / c[3] - 1) > 1e-4:
                f = f * (p / c[3])
        out.append({'iso': g['days'][i], 'o': _r2(c[0] * f), 'h': _r2(c[1] * f), 'l': _r2(c[2] * f), 'c': _r2(c[3] * f),
                    'v': _vol(c[4], f)})
        newer = i
    return out[::-1]


def _js_num(v):
    """Code.gs ab_toNum_ (JavaScript Number())."""
    if v is None or v == '':
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip() or 0)
    except ValueError:
        return float('nan')


def h60_older(h30, g, allowed):
    """{SYM: [older candles]} - Code.gs ab_h60Older_. h30 = History30 rows (header first)."""
    hdr = list(h30[0]) if h30 else []
    di = hdr.index('Date') if 'Date' in hdr else -1
    newest, first, has = '', {}, set()
    for r in h30[1:]:
        sym = str(r[0] or '').strip().upper() if r else ''
        if not sym:
            continue
        has.add(sym)
        k = h60_iso(r[di]) if 0 <= di < len(r) else ''
        if not k:
            continue
        newest = max(newest, k)
        if sym not in first or k < first[sym][0]:
            first[sym] = (k, _js_num(r[5] if len(r) > 5 else None))
    at = {d: i for i, d in enumerate(g['days'])}
    last = next((i for i in range(len(g['days']) - 1, -1, -1) if not newest or g['days'][i] <= newest), -1)

    def run_from(e):
        k = e
        while k > 0 and g['linked'][k]:
            k -= 1
        return g['days'][k]

    out = {}
    for sym in g['syms']:
        own = first.get(sym)
        if not own and sym in has:
            continue
        if not own and sym not in allowed:
            continue
        if (own[0] not in at) if own else last == -1:
            continue
        nse = h60_adjust(g, sym)
        if not own:
            frm, to = run_from(last), g['days'][last]
            lst = [c for c in nse if frm <= c['iso'] <= to]
            if lst:
                out[sym] = lst
            continue
        seam = next((c for c in nse if c['iso'] == own[0]), None)
        if not seam:
            continue
        frm = run_from(at[own[0]])
        older = [c for c in nse if frm <= c['iso'] < own[0]]
        if not older:
            continue
        if own[1] > 0:
            r = own[1] / seam['c']
            if not (0.02 < r < 50):
                continue
            if abs(r - 1) > 0.005:
                older = [{'iso': c['iso'], 'o': _r2(c['o'] * r), 'h': _r2(c['h'] * r), 'l': _r2(c['l'] * r), 'c': _r2(c['c'] * r),
                          'v': _vol(c['v'], r)} for c in older]
        out[sym] = older
    return out


def history60_data(wb):
    """window.SB_HISTORY60 - Code.gs ab_site_history60_ (the site feed part 'history60')."""
    out = {'source': 'History60 tab (NSE bhavcopy via your PC)', 'timeframe': '1d', 'stored': None,
           'dates': [], 'cols': ['d', 'o', 'h', 'l', 'c', 'v'], 'bars': {}}
    if 'History60' not in wb.sheetnames:
        return out
    g = h60_grid([list(r) for r in wb['History60'].iter_rows(values_only=True)])
    if not g['days']:
        return out
    out['stored'] = {'days': len(g['days']), 'first': g['days'][0], 'last': g['days'][-1]}
    h30 = [list(r) for r in wb['History30'].iter_rows(values_only=True)] if 'History30' in wb.sheetnames else []
    allowed = set()
    if 'PatternData' in wb.sheetnames:
        for r in list(wb['PatternData'].iter_rows(values_only=True))[1:]:
            sym = str(r[3]).strip().upper() if len(r) > 3 and r[3] is not None else ''
            if sym:
                allowed.add(sym)
    older = h60_older(h30, g, allowed)
    out['dates'] = sorted({c['iso'] for lst in older.values() for c in lst})
    at = {d: i for i, d in enumerate(out['dates'])}
    for sym, lst in older.items():
        out['bars'][sym] = [[at[c['iso']], c['o'], c['h'], c['l'], c['c'], None if math.isnan(c['v']) else c['v']] for c in lst]
    return out


def export_history60(wb):
    out = history60_data(wb)
    path = os.path.join(os.path.dirname(OUT), 'history60.js')
    js = json.dumps(out, separators=(',', ':'))
    with open(path, 'w', encoding='utf-8') as f:
        f.write('// Generated by tools/export_xlsx.py - do not edit by hand\nwindow.SB_HISTORY60 = ' + js + ';\n')
    print('wrote', os.path.normpath(path), f'{len(js) / 1024:.0f} KB,', len(out['bars']), 'symbols,', len(out['dates']), 'older dates')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    export(sys.argv[1])
